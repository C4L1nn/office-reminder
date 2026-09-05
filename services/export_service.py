"""Turn what is on screen into a file the office can send or print.

Two formats, for two different uses:

* **.xlsx** — the list as data, for sorting and filtering outside the app.
* **.pdf**  — the list as a page, for printing or e-mailing to a client.

Both export exactly the rows the user is looking at. Exporting "everything"
while the screen shows a filtered list would quietly hand someone the wrong
document.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from datetime import date, datetime
from html import escape
from pathlib import Path

from services.formatting import (
    category_label,
    recurrence_label,
    source_label,
    title_with_plate,
)

logger = logging.getLogger("office_reminder.export")

COLUMNS = ("Tarih", "Şirket", "Yükümlülük", "Dönem / Kategori", "Kaynak", "Durum")

#: Column widths in characters, for the spreadsheet.
_WIDTHS = (14, 34, 52, 30, 10, 16)


class ExportError(RuntimeError):
    """Raised with a message meant for the user, never a traceback."""


@dataclass(slots=True, frozen=True)
class ExportRow:
    due_date: date
    company: str
    title: str
    detail: str
    source: str
    status: str

    def as_tuple(self) -> tuple:
        return (self.due_date, self.company, self.title, self.detail, self.source, self.status)


def status_text(item, today: date, completed: bool = False) -> str:
    """Plain-language status, the same wording the table shows."""
    if completed:
        return "Tamamlandı"
    days = (item.due_date - today).days
    if days < 0:
        return f"{abs(days)} gün gecikti"
    if days == 0:
        return "Bugün"
    if days == 1:
        return "Yarın"
    return f"{days} gün"


def rows_from_items(items, today: date | None = None, completed: bool = False) -> list[ExportRow]:
    """Flatten due items into export rows, in the order they were given."""
    today = today or date.today()
    rows: list[ExportRow] = []
    for item in items:
        if item.source_kind == "MANUAL":
            detail = category_label(item.category)
            repeat = recurrence_label(item.recurrence_kind)
            if repeat != "—":
                detail = f"{detail} · {repeat}"
        else:
            detail = f"{source_label(item.source_label)} · {item.period_label or '—'}"
        rows.append(
            ExportRow(
                due_date=item.due_date,
                company=item.company_name or "Genel",
                title=title_with_plate(item.title, item.plate),
                detail=detail,
                source="Resmî" if item.source_kind == "OFFICIAL" else "Manuel",
                status=status_text(item, today, completed),
            )
        )
    return rows


# --------------------------------------------------------------------- xlsx
def write_xlsx(rows: list[ExportRow], path: Path, title: str) -> Path:
    """Write a real spreadsheet: bold header, frozen top row, autofilter."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - packaging guard
        raise ExportError(
            "Excel dışa aktarma bileşeni bulunamadı. Uygulamayı yeniden kurmayı deneyin."
        ) from exc

    book = Workbook()
    sheet = book.active
    # Excel refuses sheet names over 31 characters or containing []:*?/\
    sheet.title = "Liste"

    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        sheet.append(list(row.as_tuple()))

    for index, width in enumerate(_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet["A"][1:]:
        cell.number_format = "DD.MM.YYYY"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}"

    try:
        book.save(path)
    except OSError as exc:
        raise ExportError(f"Dosya yazılamadı: {exc.strerror or exc}") from exc
    logger.info("xlsx export: %s rows -> %s", len(rows), path)
    return path


# ---------------------------------------------------------------------- pdf
def _table_html(rows: list[ExportRow], title: str) -> str:
    """The printable page. Built as HTML because QTextDocument lays out tables
    far better than hand-placed QPainter text."""
    header = "".join(f"<th>{escape(column)}</th>" for column in COLUMNS)
    body = []
    for row in rows:
        body.append(
            "<tr>"
            f"<td class='date'>{row.due_date.strftime('%d.%m.%Y')}</td>"
            f"<td>{escape(row.company)}</td>"
            f"<td>{escape(row.title)}</td>"
            f"<td>{escape(row.detail)}</td>"
            f"<td>{escape(row.source)}</td>"
            f"<td>{escape(row.status)}</td>"
            "</tr>"
        )
    stamp = datetime.now().strftime("%d.%m.%Y %H:%M")
    return f"""
    <html><head><meta charset="utf-8"><style>
      body {{ font-family: "Segoe UI", sans-serif; font-size: 9pt; color: #1c222c; }}
      h1 {{ font-size: 15pt; margin: 0 0 2px 0; }}
      .meta {{ color: #667085; font-size: 8pt; margin-bottom: 10px; }}
      table {{ border-collapse: collapse; width: 100%; }}
      th {{ background: #eef1f6; text-align: left; font-size: 8pt;
            padding: 5px 6px; border-bottom: 1px solid #c4ccd8; }}
      td {{ padding: 4px 6px; border-bottom: 1px solid #e6e9ef; }}
      td.date {{ white-space: nowrap; }}
    </style></head><body>
      <h1>{escape(title)}</h1>
      <div class="meta">{len(rows)} kayıt · {escape(stamp)} · Office Reminder</div>
      <table><thead><tr>{header}</tr></thead><tbody>{''.join(body)}</tbody></table>
    </body></html>
    """


def write_pdf(rows: list[ExportRow], path: Path, title: str) -> Path:
    """Render the list to an A4 PDF."""
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument

    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageMargins(QMarginsF(12, 12, 12, 14), QPageLayout.Unit.Millimeter)
    writer.setResolution(150)
    writer.setTitle(title)

    document = QTextDocument()
    document.setHtml(_table_html(rows, title))
    document.setPageSize(writer.pageLayout().paintRectPixels(writer.resolution()).size().toSizeF())
    try:
        document.print_(writer)
    except Exception as exc:  # pragma: no cover - Qt failure path
        raise ExportError(f"PDF oluşturulamadı: {exc}") from exc
    logger.info("pdf export: %s rows -> %s", len(rows), path)
    return path


# -------------------------------------------------------------------- notes
def write_board_pdf(render, source_size, path: Path, title: str) -> Path:
    """Print a sticky-note board exactly as it is arranged on screen.

    `render(painter, target_rect)` does the drawing, so this function stays
    free of any knowledge of the canvas. The page is turned on its side when
    the board is wider than it is tall, and the board is scaled to fit rather
    than split, so one board is always one page.
    """
    from PySide6.QtCore import QMarginsF, QRectF
    from PySide6.QtGui import QPageLayout, QPageSize, QPainter, QPdfWriter

    width, height = float(source_size[0]), float(source_size[1])
    if width <= 0 or height <= 0:
        raise ExportError("Bu sayfada dışa aktarılacak not yok.")

    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageOrientation(
        QPageLayout.Orientation.Landscape
        if width > height
        else QPageLayout.Orientation.Portrait
    )
    writer.setPageMargins(QMarginsF(10, 10, 10, 10), QPageLayout.Unit.Millimeter)
    writer.setResolution(300)
    writer.setTitle(title)

    page = writer.pageLayout().paintRectPixels(writer.resolution())
    painter = QPainter(writer)
    try:
        # Keep the board's proportions; a stretched note is not "as I left it".
        scale = min(page.width() / width, page.height() / height)
        drawn = QRectF(0, 0, width * scale, height * scale)
        drawn.moveCenter(QRectF(page).center())
        render(painter, drawn)
    finally:
        painter.end()
    logger.info("board pdf export: %sx%s -> %s", int(width), int(height), path)
    return path


def default_name(prefix: str, suffix: str) -> str:
    """A filename that sorts by date and never collides across days."""
    return f"{prefix}_{date.today():%Y-%m-%d}.{suffix}"


__all__ = [
    "COLUMNS",
    "ExportError",
    "ExportRow",
    "default_name",
    "rows_from_items",
    "status_text",
    "write_board_pdf",
    "write_pdf",
    "write_xlsx",
]
